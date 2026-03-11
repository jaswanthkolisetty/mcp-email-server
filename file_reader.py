import io
import base64

# Maps file extensions to their MIME types
_EXT_TO_MIME = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "csv": "text/csv",
    "txt": "text/plain",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "gif": "image/gif",
    "webp": "image/webp",
}


def get_content_type(filename: str) -> str:
    """Derive MIME type from a filename's extension."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    return _EXT_TO_MIME.get(ext, "application/octet-stream")


def extract_text(content: bytes, filename: str, content_type: str = "") -> str:
    """
    Extract human-readable text (or a data URI for images) from file bytes.
    Detects format using the filename extension, with content_type as fallback.
    """
    if not content_type:
        content_type = get_content_type(filename)
    name = filename.lower()

    if name.endswith(".pdf") or "pdf" in content_type:
        return _read_pdf(content)
    elif name.endswith((".xlsx", ".xls")) or "spreadsheet" in content_type or "excel" in content_type:
        return _read_excel(content)
    elif name.endswith(".docx") or "wordprocessingml" in content_type:
        return _read_docx(content)
    elif name.endswith(".csv") or "csv" in content_type:
        return content.decode("utf-8", errors="replace")
    elif name.endswith(".txt") or content_type.startswith("text/"):
        return content.decode("utf-8", errors="replace")
    elif name.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")) or content_type.startswith("image/"):
        # Return as a base64 data URI so the AI model can view it
        ext = name.rsplit(".", 1)[-1] if "." in name else "png"
        b64 = base64.b64encode(content).decode()
        return f"data:image/{ext};base64,{b64}"
    else:
        return f"[Binary file: {filename} ({len(content) // 1024} KB) — cannot extract text]"


def _read_pdf(content: bytes) -> str:
    """Extract text from all pages of a PDF, separated by page break markers."""
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n--- Page Break ---\n\n".join(pages).strip()
    except ImportError:
        return "[PyPDF2 not installed]"
    except Exception as e:
        return f"[Error reading PDF: {e}]"


def _read_excel(content: bytes) -> str:
    """Extract all cell values from every sheet, tab-separated per row."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        output = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            output.append(f"=== Sheet: {sheet_name} ===")
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    output.append("\t".join("" if c is None else str(c) for c in row))
        return "\n".join(output)
    except ImportError:
        return "[openpyxl not installed]"
    except Exception as e:
        return f"[Error reading Excel: {e}]"


def _read_docx(content: bytes) -> str:
    """Extract paragraph text from a Word document, skipping empty paragraphs."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return "[python-docx not installed]"
    except Exception as e:
        return f"[Error reading Word doc: {e}]"
